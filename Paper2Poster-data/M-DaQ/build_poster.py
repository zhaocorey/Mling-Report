#!/usr/bin/env python3
"""Build M-DaQ poster for SIGIR 2026. A0 Portrait, 3-column. Low memory."""
import os
from pptx import Presentation
from pptx.util import Pt, Cm, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# === Layout ===
W_IN, H_IN = 33.1, 46.8
MARGIN = Cm(3)
COL_GAP = Cm(2)

# Colors (as hex strings for XML fallback)
C_PRIMARY   = RGBColor(0x1B, 0x3A, 0x5C)
C_ACCENT    = RGBColor(0xE8, 0x79, 0x2B)
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
C_TEXT      = RGBColor(0x22, 0x22, 0x22)
C_GRAY      = RGBColor(0x66, 0x66, 0x66)
C_LGRAY     = RGBColor(0x99, 0x99, 0x99)
C_CARD      = RGBColor(0xF0, 0xF4, 0xF8)
C_LIGHTBLUE = RGBColor(0xCC, 0xDD, 0xEE)

HEX = {
    C_PRIMARY: "1B3A5C", C_ACCENT: "E8792B", C_WHITE: "FFFFFF",
    C_TEXT: "222222", C_GRAY: "666666", C_LGRAY: "999999",
    C_CARD: "F0F4F8", C_LIGHTBLUE: "CCDDEE",
}

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(OUT_DIR, "_output")
LOGO_DIR = os.path.join(OUT_DIR, "logos")

prs = Presentation()
prs.slide_width = Inches(W_IN)
prs.slide_height = Inches(H_IN)
slide = prs.slides.add_slide(prs.slide_layouts[6])

def _set_color(run, color):
    """Set font color via XML (workaround for pptx 1.0.2 bug)."""
    hex_val = HEX.get(color, "222222")
    rPr = run._r.get_or_add_rPr()
    # Remove existing solidFill if any
    for old in rPr.findall(qn('a:solidFill')):
        rPr.remove(old)
    sf = etree.SubElement(rPr, qn('a:solidFill'))
    sc = etree.SubElement(sf, qn('a:srgbClr'))
    sc.set('val', hex_val)

def tb(left, top, w, h, text, sz=24, bold=False, color=C_TEXT, align=PP_ALIGN.LEFT, ls=1.25):
    box = slide.shapes.add_textbox(left, top, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sz * 0.25)
        if ls:
            p.line_spacing = ls
        r = p.add_run()
        r.text = line
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.name = "Arial"
        _set_color(r, color)
    return box

def rect(l, t, w, h, c):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = c
    s.line.fill.background()
    return s

def img(path, l, t, w=None):
    if os.path.exists(path):
        if w:
            slide.shapes.add_picture(path, l, t, width=w)
        else:
            slide.shapes.add_picture(path, l, t)

def header(x, y, w, title):
    tb(x, y, w, Cm(2), title, sz=30, bold=True, color=C_PRIMARY)
    rect(x, y + Cm(2), w, Pt(3), C_ACCENT)
    return y + Cm(2.8)

def bullets(x, y, w, items, sz=19):
    text = "\n".join(f"\u2022 {i}" for i in items)
    h = Cm(len(items) * 1.1 + 0.5)
    tb(x, y, w, h, text, sz=sz, ls=1.2)
    return y + h

def tbl(x, y, w, heads, rows, sz=16):
    nc = len(heads)
    cw = w // nc
    rh = Cm(1.1)
    for ci, h in enumerate(heads):
        rect(x + ci * cw, y, cw, rh, C_PRIMARY)
        tb(x + ci * cw + Pt(4), y + Pt(1), cw - Pt(8), rh, h,
           sz=sz, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    y += rh
    for ri, row in enumerate(rows):
        bg = C_CARD if ri % 2 == 0 else C_WHITE
        for ci, cell in enumerate(row):
            rect(x + ci * cw, y, cw, rh, bg)
            tb(x + ci * cw + Pt(4), y + Pt(1), cw - Pt(8), rh, cell,
               sz=sz, color=C_TEXT, align=PP_ALIGN.CENTER)
        y += rh
    return y + Cm(0.3)

# ============================================================
# CONTENT
# ============================================================
cl, ct = MARGIN, MARGIN
cw_full = prs.slide_width - 2 * MARGIN
ch_full = prs.slide_height - 2 * MARGIN

# ── Title Bar ──
TH = Cm(8)
rect(cl, ct, cw_full, TH, C_PRIMARY)
tb(cl + Cm(1), ct + Cm(0.8), cw_full - Cm(2), Cm(4.5),
   "M-DaQ: Retrieving Samples with Multilingual Diversity\nand Quality for Instruction Fine-Tuning Datasets",
   sz=58, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, ls=1.15)
tb(cl + Cm(1), ct + Cm(5.2), cw_full - Cm(2), Cm(1.3),
   "Chunguang Zhao\u00b9 \u00b7 Yilun Liu\u00b9 \u00b7 Pufan Zeng\u00b2 \u00b7 Yuanchang Luo\u00b9 \u00b7 Shimin Tao\u00b9 \u00b7 Minggui He\u00b9 \u00b7 Weibin Meng\u00b9 \u00b7 Song Xu\u00b2 \u00b7 Chen Liu\u00b9 \u00b7 Hongxia Ma\u00b9 \u00b7 Li Zhang\u00b9 \u00b7 Boxing Chen\u00b9 \u00b7 Daimeng Wei\u00b9",
   sz=18, color=C_WHITE, align=PP_ALIGN.CENTER)
tb(cl + Cm(1), ct + Cm(6.5), cw_full - Cm(2), Cm(1),
   "\u00b9 Huawei Technologies Ltd.    \u00b2 University of Science and Technology of China",
   sz=15, color=C_LIGHTBLUE, align=PP_ALIGN.CENTER)

# Logos
img(os.path.join(LOGO_DIR, "huawei_logo.png"), cl + Cm(0.5), ct + Cm(1), Cm(5))
img(os.path.join(LOGO_DIR, "ustc_logo.png"), cl + cw_full - Cm(6), ct + Cm(1), Cm(5))
sigir = "/home/admin/.openclaw/agents/zong/workspace/Paper2Poster/logo_store/conferences/sigir.png"
if os.path.exists(sigir):
    img(sigir, cl + cw_full // 2 - Cm(2), ct + Cm(0.8), Cm(4))

# ── Columns ──
col_top = ct + TH + Cm(1)
uw = cw_full - 2 * COL_GAP
col_w = [int(uw * 0.30), int(uw * 0.40), int(uw * 0.30)]
col_x = [cl, cl + col_w[0] + COL_GAP, cl + col_w[0] + COL_GAP + col_w[1] + COL_GAP]

# ═══ COL 1: Motivation + Setup + SAH ═══
x, y = col_x[0], col_top
y = header(x, y, col_w[0], "Motivation")
y = bullets(x, y, col_w[0], [
    "Multilingual IFT data is scarce & skewed toward English",
    "Llama-3 IFT: only 3.01% multilingual samples",
    "No language-agnostic quality scoring method",
    "SAH unverified in multilingual contexts",
])
y += Cm(0.4)
y = tbl(x, y, col_w[0], ["Challenge", "M-DaQ Solution"], [
    ["No multilingual quality scoring", "QSM (triplet loss)"],
    ["Diversity limited to English", "DAS (MMR-inspired)"],
    ["SAH unverified multilingually", "1K\u219252K, 8 langs"],
], sz=15)
y += Cm(0.8)

y = header(x, y, col_w[0], "Experimental Setup")
y = bullets(x, y, col_w[0], [
    "Base model: Llama-3-8B",
    "IFT data: Alpaca-52K \u00d7 18 languages",
    "Training: 3 epochs, lr 5\u00d710\u207b\u2075",
    "Judges: LLM-as-Judge + 7 native experts",
    "Benchmarks: Alpaca-Eval + MT-Bench",
])
y += Cm(0.8)

y = header(x, y, col_w[0], "SAH Validation")
img(os.path.join(FIG_DIR, "fig_p4_0.png"), x + Cm(0.3), y, col_w[0] - Cm(0.6))
y += Cm(6.5)
tb(x, y, col_w[0], Cm(3),
   "\u2022 1K curated > 52K unfiltered \u2014 SAH holds!\n\u2022 Diminishing returns beyond quality threshold\n\u2022 Language-dependent sensitivity observed",
   sz=17, ls=1.2)

# ═══ COL 2: Method + Results (HERO) ═══
x, y = col_x[1], col_top
y = header(x, y, col_w[1], "Method: M-DaQ Framework")
img(os.path.join(FIG_DIR, "fig_p1_0.png"), x + Cm(0.5), y, col_w[1] - Cm(1))
y += Cm(7.5)

tb(x, y, col_w[1], Cm(1.5), "Stage 1: Quality Scoring Model (QSM)", sz=21, bold=True, color=C_ACCENT)
y += Cm(1.8)
y = bullets(x, y, col_w[1], [
    "Fine-tuned on ~2.3K expert-revised samples \u00d7 18 languages",
    "Triplet loss: instruction \u2192 positive vs. negative embedding",
    "Language-agnostic quality signal from embedding space",
])
y += Cm(0.3)

tb(x, y, col_w[1], Cm(1.5), "Stage 2: Diversity-Aware Selection (DAS)", sz=21, bold=True, color=C_ACCENT)
y += Cm(1.8)
y = bullets(x, y, col_w[1], [
    "MMR-inspired: Quality \u2248 Relevance, Diversity \u2248 Novelty",
    "Select top-n by QSM \u2192 greedily augment from clusters",
    "Ratio n_quality : n_diversity = 6:1 (tuned)",
    "Complexity: O(n\u00b2) \u2192 O(n log n)",
])
y += Cm(0.8)

y = header(x, y, col_w[1], "Cross-Lingual Results (18 Languages)")
img(os.path.join(FIG_DIR, "fig_p1_1.png"), x + Cm(0.5), y, col_w[1] - Cm(1))
y += Cm(7.5)

rect(x + Cm(0.5), y, col_w[1] - Cm(1), Cm(3), C_PRIMARY)
tb(x + Cm(1), y + Cm(0.3), col_w[1] - Cm(2), Cm(2.5),
   "Avg. Win Rate:  Alpaca-Eval 60.2%  |  MT-Bench R1 62.6%  |  R2 62.9%\nGains more pronounced in low-resource languages (Tagalog, Malay)",
   sz=18, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER, ls=1.3)

# ═══ COL 3: Key Results + Conclusion ═══
x, y = col_x[2], col_top
y = header(x, y, col_w[2], "Key Results")

tb(x, y, col_w[2], Cm(1.5), "LLM-as-Judge: M-DaQ vs. Vanilla", sz=19, bold=True, color=C_ACCENT)
y += Cm(1.6)
y = tbl(x, y, col_w[2], ["Benchmark", "Win Rate"], [
    ["Alpaca-Eval", "60.2%"],
    ["MT-Bench R1", "62.6%"],
    ["MT-Bench R2", "62.9%"],
], sz=16)
y += Cm(0.4)

tb(x, y, col_w[2], Cm(1.5), "Human Eval (6 Langs, 58 Person-Hours)", sz=19, bold=True, color=C_ACCENT)
y += Cm(1.6)
y = tbl(x, y, col_w[2], ["Language", "AE", "MT-R1", "MT-R2"], [
    ["Japanese", "86%", "84%", "80%"],
    ["Korean", "94%", "94%", "94%"],
    ["Russian", "70%", "86%", "84%"],
    ["Portuguese", "80%", "78%", "84%"],
    ["Greek", "58%", "76%", "82%"],
    ["French", "80%", "92%", "88%"],
    ["Average", "78.0%", "85.0%", "85.3%"],
], sz=15)
y += Cm(0.4)

# Corsica
y = header(x, y, col_w[2], "Cultural Localization")
img(os.path.join(FIG_DIR, "fig_p4_1.png"), x + Cm(0.3), y, col_w[2] - Cm(0.6))
y += Cm(5.5)
tb(x, y, col_w[2], Cm(2),
   "M-DaQ incorporates local knowledge (Piana, figatelli)\nvs. generic baseline responses",
   sz=15, color=C_GRAY, ls=1.15)
y += Cm(2.2)

# Conclusion
y = header(x, y, col_w[2], "Conclusion")
y = bullets(x, y, col_w[2], [
    "QSM + DAS \u2192 compact, high-fidelity multilingual IFT",
    "60%+ avg. win rate across 18 languages",
    "Human eval: cultural relevance & instruction-following",
    "First SAH validation in multilingual settings",
    "Few thousand quality samples > 52K unfiltered",
])
y += Cm(0.8)

# Contact card
rect(x + Cm(0.3), y, col_w[2] - Cm(0.6), Cm(3.5), C_CARD)
tb(x + Cm(0.8), y + Cm(0.3), col_w[2] - Cm(1.6), Cm(3),
   "\U0001f4ce github.com/zhaocorey/M-DaQ\n\U0001f4ce arXiv:2509.15549\n\U0001f4e7 liuyilun3@huawei.com",
   sz=17, color=C_PRIMARY, ls=1.3)

# ── Footer ──
tb(cl, prs.slide_height - MARGIN + Cm(0.5), cw_full, Cm(1.5),
   "SIGIR 2026  \u00b7  Melbourne, VIC, Australia  \u00b7  July 20\u201324, 2026",
   sz=15, color=C_LGRAY, align=PP_ALIGN.CENTER)

# ── Save ──
out = os.path.join(OUT_DIR, "_output", "M-DaQ_poster_v2.pptx")
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print(f"\u2705 Poster saved: {out}  ({os.path.getsize(out)//1024} KB)")
