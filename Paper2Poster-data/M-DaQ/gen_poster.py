#!/usr/bin/env python3
"""Simplified poster generator for M-DaQ @ SIGIR 2026.
Extracts figures from PDF + generates A0 portrait poster PPTX.
"""
import os
import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image
import io

# ── Config ──
PAPER_PDF = "paper.pdf"
OUTPUT_DIR = "_output"
POSTER_W_INCHES = 33.1  # A0 width
POSTER_H_INCHES = 46.8  # A0 height

# Colors
C_DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
C_ACCENT_ORANGE = RGBColor(0xE8, 0x79, 0x2B)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK = RGBColor(0x22, 0x22, 0x22)
C_GRAY = RGBColor(0x66, 0x66, 0x66)
C_LIGHT_BG = RGBColor(0xF0, 0xF4, 0xF8)

os.makedirs(OUTPUT_DIR, exist_ok=True)

# ── Step 1: Extract images from PDF ──
print("📄 Extracting figures from PDF...")
doc = fitz.open(PAPER_PDF)
figures = []
page_img_counts = {}
for page_idx in range(len(doc)):
    page = doc[page_idx]
    image_list = page.get_images(full=True)
    page_img_counts[page_idx] = 0
    for img_idx, img_info in enumerate(image_list):
        xref = img_info[0]
        base_image = doc.extract_image(xref)
        if base_image and base_image["width"] > 200 and base_image["height"] > 150:
            img_data = base_image["image"]
            img_path = os.path.join(OUTPUT_DIR, f"fig_p{page_idx}_{img_idx}.png")
            with open(img_path, "wb") as f:
                f.write(img_data)
            figures.append({
                "path": img_path,
                "page": page_idx,
                "img_idx": img_idx,
                "w": base_image["width"],
                "h": base_image["height"],
                "ext": base_image["ext"]
            })
            print(f"  ✅ Page {page_idx}, image {img_idx}: {base_image['width']}x{base_image['height']}")

print(f"  Total: {len(figures)} figures extracted")
doc.close()

# ── Step 2: Create poster PPTX ──
print("\n🎨 Generating poster PPTX...")
prs = Presentation()
prs.slide_width = Inches(POSTER_W_INCHES)
prs.slide_height = Inches(POSTER_H_INCHES)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

from pptx.oxml.ns import qn
from lxml import etree

def set_run_color(run, rgb_color):
    """Set run color via XML (python-pptx 1.0.2 compat)."""
    rPr = run._r.get_or_add_rPr()
    # Remove existing solidFill
    for sf in rPr.findall(qn('a:solidFill')):
        rPr.remove(sf)
    solidFill = etree.SubElement(rPr, qn('a:solidFill'))
    srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgbClr.set('val', str(rgb_color))

def add_textbox(slide, left, top, width, height, text, font_size=24,
                bold=False, color=C_BLACK, alignment=PP_ALIGN.LEFT,
                font_name="Arial", fill_color=None):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    set_run_color(run, color)
    run.font.name = font_name
    p.alignment = alignment
    if fill_color:
        txBox.fill.solid()
        txBox.fill.fore_color.rgb = fill_color
    return txBox

def add_bullets(slide, left, top, width, height, items, font_size=22,
                color=C_BLACK, spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        run = p.add_run()
        run.text = f"• {item}"
        run.font.size = Pt(font_size)
        set_run_color(run, color)
        run.font.name = "Arial"
        p.space_after = spacing
    return txBox

def add_image_safe(slide, img_path, left, top, width=None, height=None):
    if os.path.exists(img_path):
        kwargs = {"image_file": img_path,
                  "left": Inches(left), "top": Inches(top)}
        if width:
            kwargs["width"] = Inches(width)
        if height:
            kwargs["height"] = Inches(height)
        slide.shapes.add_picture(**kwargs)
        return True
    return False

# ── Margins & Layout ──
MARGIN = 0.8
COL_GAP = 0.5
CONTENT_W = POSTER_W_INCHES - 2 * MARGIN
COL_W = (CONTENT_W - 2 * COL_GAP) / 3
COL1_X = MARGIN
COL2_X = MARGIN + COL_W + COL_GAP
COL3_X = MARGIN + 2 * (COL_W + COL_GAP)

# ── TITLE BAR ──
title_h = 2.8
title_box = slide.shapes.add_shape(
    1, Inches(0), Inches(0), Inches(POSTER_W_INCHES), Inches(title_h))  # Rectangle
title_box.fill.solid()
title_box.fill.fore_color.rgb = C_DARK_BLUE
title_box.line.fill.background()

add_textbox(slide, 0.8, 0.3, POSTER_W_INCHES - 1.6, 1.5,
            "M-DaQ: Retrieving Samples with Multilingual Diversity and Quality\nfor Instruction Fine-Tuning Datasets",
            font_size=44, bold=True, color=C_WHITE, alignment=PP_ALIGN.CENTER)

add_textbox(slide, 0.8, 1.9, POSTER_W_INCHES - 1.6, 0.5,
            "Chunguang Zhao¹ · Yilun Liu¹ · Pufan Zeng² · Yuanchang Luo¹ · Shimin Tao¹ · Minggui He¹ · Weibin Meng¹ · Song Xu² · Chen Liu¹ · Hongxia Ma¹ · Li Zhang¹ · Boxing Chen¹ · Daimeng Wei¹",
            font_size=18, color=C_WHITE, alignment=PP_ALIGN.CENTER)

add_textbox(slide, 0.8, 2.3, POSTER_W_INCHES - 1.6, 0.4,
            "¹ Huawei Technologies Ltd.    ² University of Science and Technology of China  |  SIGIR 2026  |  github.com/zhaocorey/M-DaQ",
            font_size=16, color=RGBColor(0xBB, 0xCC, 0xDD), alignment=PP_ALIGN.CENTER)

Y_START = title_h + 0.4

# ── SECTION HEADER helper ──
def section_header(slide, x, y, text, number=None):
    label = f"{number}  {text}" if number else text
    add_textbox(slide, x, y, COL_W, 0.5, label,
                font_size=28, bold=True, color=C_DARK_BLUE)
    # Underline
    line = slide.shapes.add_shape(1, Inches(x), Inches(y + 0.5),
                                   Inches(COL_W), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = C_ACCENT_ORANGE
    line.line.fill.background()
    return y + 0.7

# ═══════════════════════════════════════
# COLUMN 1 (Left)
# ═══════════════════════════════════════
y1 = Y_START

# 1. MOTIVATION
y1 = section_header(slide, COL1_X, y1, "MOTIVATION", "1")
add_bullets(slide, COL1_X, y1, COL_W, 3.0, [
    "Multilingual IFT data is scarce, skewed toward English, and lacks systematic curation",
    "Llama-3 IFT dataset: only 3.01% multilingual samples",
    "No language-agnostic quality scoring method for multilingual data",
    "Data diversity studied only in English settings",
    "Superficial Alignment Hypothesis (SAH) unverified in multilingual contexts",
], font_size=20)
y1 += 3.2

# Three Challenges table
add_textbox(slide, COL1_X, y1, COL_W, 0.4, "Three Challenges Addressed:",
            font_size=20, bold=True, color=C_DARK_BLUE)
y1 += 0.5

challenges = [
    ("C1", "No extensible quality scoring", "→ QSM with triplet loss"),
    ("C2", "Diversity selection English-only", "→ DAS inspired by MMR"),
    ("C3", "SAH unverified multilingually", "→ 1K→52K scale study, 8 langs"),
]
for cid, challenge, solution in challenges:
    add_textbox(slide, COL1_X, y1, 0.6, 0.35, cid,
                font_size=18, bold=True, color=C_ACCENT_ORANGE)
    add_textbox(slide, COL1_X + 0.6, y1, COL_W - 0.6, 0.35,
                f"{challenge}  {solution}", font_size=18, color=C_BLACK)
    y1 += 0.45

y1 += 0.3

# 2. METHOD: QSM
y1 = section_header(slide, COL1_X, y1, "METHOD", "2")

add_textbox(slide, COL1_X, y1, COL_W, 0.4, "Stage 1: Quality Scoring Model (QSM)",
            font_size=20, bold=True, color=C_ACCENT_ORANGE)
y1 += 0.45
add_bullets(slide, COL1_X, y1, COL_W, 1.8, [
    "Fine-tuned on ~2.3K expert-revised samples × 18 languages",
    "Triplet loss: instruction → positive (expert) vs negative (original/MT)",
    "Language-agnostic quality signal from embedding space",
], font_size=18)
y1 += 1.8

add_textbox(slide, COL1_X, y1, COL_W, 0.4, "Stage 2: Diversity-Aware Selection (DAS)",
            font_size=20, bold=True, color=C_ACCENT_ORANGE)
y1 += 0.45
add_bullets(slide, COL1_X, y1, COL_W, 2.2, [
    "MMR-inspired: Quality ≈ Relevance, Diversity ≈ Novelty",
    "Two-stage pipeline: O(n²) → O(n log n)",
    "  Stage 1: Select top-n by QSM score (quality)",
    "  Stage 2: Greedily augment from uncovered clusters (diversity)",
    "Ratio n_quality : n_diversity = 6:1 (empirically tuned)",
], font_size=18)
y1 += 2.3

# ── Figure mapping (identified by page + dimensions) ──
# fig_p1_0 (2491x1266) = Figure 1b: 18-lang win rate bars (HERO)
# fig_p1_1 (2400x1864) = Figure 2: M-DaQ pipeline diagram
# fig_p2_0 (456x394)   = Algorithm 1 / small diagram
# fig_p4_0 (828x414)   = Figure 3: Corsica cultural example
# fig_p4_1 (1200x840)  = Figure 4: SAH validation line chart

fig_map = {}
for f in figures:
    key = f"fig_p{f['page']}_{f['img_idx']}"
    fig_map[key] = f

winrate_fig = fig_map.get("fig_p1_0")     # 18-lang win rates (hero)
pipeline_fig = fig_map.get("fig_p1_1")    # M-DaQ pipeline
sah_fig = fig_map.get("fig_p4_1")         # SAH line chart
corsica_fig = fig_map.get("fig_p4_0")     # Corsica cultural example

if pipeline_fig:
    add_image_safe(slide, pipeline_fig["path"], COL1_X, y1, width=COL_W)
    add_textbox(slide, COL1_X, y1 + 3.2, COL_W, 0.4,
                "Figure: M-DaQ two-stage pipeline (QSM + DAS)",
                font_size=16, color=C_GRAY, alignment=PP_ALIGN.CENTER)
    y1 += 3.8

# ═══════════════════════════════════════
# COLUMN 2 (Center - Hero)
# ═══════════════════════════════════════
y2 = Y_START

# 3. RESULTS
y2 = section_header(slide, COL2_X, y2, "KEY RESULTS", "3")

# Win rate highlight box
highlight_box = slide.shapes.add_shape(
    1, Inches(COL2_X), Inches(y2), Inches(COL_W), Inches(1.0))
highlight_box.fill.solid()
highlight_box.fill.fore_color.rgb = C_LIGHT_BG
highlight_box.line.fill.background()

add_textbox(slide, COL2_X + 0.2, y2 + 0.05, COL_W - 0.4, 0.4,
            "LLM-as-Judge: Avg Win Rate (M-DaQ vs Vanilla)",
            font_size=18, bold=True, color=C_DARK_BLUE)

results_text = "Alpaca-Eval: 60.2%  |  MT-Bench R1: 62.6%  |  MT-Bench R2: 62.9%"
add_textbox(slide, COL2_X + 0.2, y2 + 0.45, COL_W - 0.4, 0.4,
            results_text, font_size=18, bold=True, color=C_ACCENT_ORANGE)
y2 += 1.2

if winrate_fig:
    img_h = min(5.0, winrate_fig["h"] * COL_W / winrate_fig["w"])
    add_image_safe(slide, winrate_fig["path"], COL2_X, y2, width=COL_W)
    add_textbox(slide, COL2_X, y2 + img_h + 0.05, COL_W, 0.4,
                "Figure: Cross-lingual win rates across 18 languages (Alpaca-Eval + MT-Bench)",
                font_size=16, color=C_GRAY, alignment=PP_ALIGN.CENTER)
    y2 += img_h + 0.6
else:
    y2 += 0.3

# Key observations
add_textbox(slide, COL2_X, y2, COL_W, 0.4, "Key Observations:",
            font_size=20, bold=True, color=C_DARK_BLUE)
y2 += 0.4
add_bullets(slide, COL2_X, y2, COL_W, 2.0, [
    "Consistent improvement across all 18 languages",
    "Gains MORE pronounced in low-resource languages (Tagalog, Malay)",
    "M-DaQ effectively mitigates linguistic imbalance and data scarcity",
], font_size=18)
y2 += 2.0

# Human Eval Table
add_textbox(slide, COL2_X, y2, COL_W, 0.4, "Human Evaluation (6 Langs, 900 Samples, 58 Person-Hours):",
            font_size=18, bold=True, color=C_DARK_BLUE)
y2 += 0.5

# Simple table as text
table_data = [
    "Language      Alpaca-Eval   MT-Bench R1   MT-Bench R2",
    "Japanese      86%           84%           80%",
    "Korean        94%           94%           94%",
    "Russian       70%           86%           84%",
    "Portuguese    80%           78%           84%",
    "Greek         58%           76%           82%",
    "French        80%           92%           88%",
    "Average       78.0%         85.0%         85.3%",
]
table_box = slide.shapes.add_textbox(Inches(COL2_X), Inches(y2),
                                      Inches(COL_W), Inches(2.5))
tf = table_box.text_frame
tf.word_wrap = True
for i, row in enumerate(table_data):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    run = p.add_run()
    run.text = row
    run.font.size = Pt(16)
    run.font.name = "Courier New"
    set_run_color(run, C_DARK_BLUE if i == 0 or i == len(table_data)-1 else C_BLACK)
    run.font.bold = (i == 0 or i == len(table_data)-1)
    p.space_after = Pt(2)
y2 += 2.8

# ═══════════════════════════════════════
# COLUMN 3 (Right)
# ═══════════════════════════════════════
y3 = Y_START

# 4. SAH Validation
y3 = section_header(slide, COL3_X, y3, "SAH IN MULTILINGUAL SETTINGS", "4")

add_textbox(slide, COL3_X, y3, COL_W, 0.4,
            "First systematic investigation of the Superficial Alignment Hypothesis across languages.",
            font_size=18, color=C_BLACK)
y3 += 0.6

if sah_fig:
    img_h = min(4.0, sah_fig["h"] * COL_W / sah_fig["w"])
    add_image_safe(slide, sah_fig["path"], COL3_X, y3, width=COL_W)
    add_textbox(slide, COL3_X, y3 + img_h + 0.05, COL_W, 0.4,
                "Figure: Win rate vs IFT dataset scale (1K–52K), 8 languages",
                font_size=16, color=C_GRAY, alignment=PP_ALIGN.CENTER)
    y3 += img_h + 0.6

# SAH findings
add_textbox(slide, COL3_X, y3, COL_W, 0.4, "Key Findings:",
            font_size=20, bold=True, color=C_DARK_BLUE)
y3 += 0.4

add_bullets(slide, COL3_X, y3, COL_W, 2.0, [
    "✅ Diminishing returns: 1K curated > 52K unfiltered — SAH holds",
    "  10K: −10.1% avg win rate vs 1K baseline",
    "  52K: −6.2% avg win rate vs 1K baseline",
    "⚠️ Language-dependent sensitivity:",
    "  Arabic shows flatter curve than French",
    "  (lower cross-lingual pretraining readiness)",
    "💡 A few thousand high-quality samples suffice for effective multilingual alignment",
], font_size=18)
y3 += 2.8

# 5. Cultural Localization
y3 = section_header(slide, COL3_X, y3, "CULTURAL LOCALIZATION", "5")

if corsica_fig:
    img_h = min(3.0, corsica_fig["h"] * COL_W / corsica_fig["w"])
    add_image_safe(slide, corsica_fig["path"], COL3_X, y3, width=COL_W)
    add_textbox(slide, COL3_X, y3 + img_h + 0.05, COL_W, 0.5,
                "Example: French travel query about Corsica → M-DaQ references Piana, figatelli, local conventions",
                font_size=16, color=C_GRAY, alignment=PP_ALIGN.CENTER)
    y3 += img_h + 0.7
else:
    add_bullets(slide, COL3_X, y3, COL_W, 1.5, [
        "French travel query about Corsica:",
        "M-DaQ: references Piana, figatelli, local conventions",
        "Baseline: generic, culturally neutral response",
    ], font_size=18)
    y3 += 1.5

# 6. CONCLUSION
y3 = section_header(slide, COL3_X, y3, "CONCLUSION", "6")

add_bullets(slide, COL3_X, y3, COL_W, 2.5, [
    "M-DaQ = QSM (quality) + DAS (diversity): compact, high-fidelity multilingual IFT subsets",
    "60%+ avg win rate across 18 languages on Alpaca-Eval & MT-Bench",
    "Human eval confirms gains in cultural relevance, contextual appropriateness, instruction-following",
    "First empirical validation of SAH in multilingual settings",
    "Code: github.com/zhaocorey/M-DaQ",
], font_size=18)
y3 += 2.5

# QR / Contact
y3 += 0.3
contact_box = slide.shapes.add_shape(
    1, Inches(COL3_X), Inches(y3), Inches(COL_W), Inches(1.2))
contact_box.fill.solid()
contact_box.fill.fore_color.rgb = C_LIGHT_BG
contact_box.line.fill.background()

add_textbox(slide, COL3_X + 0.2, y3 + 0.1, COL_W - 0.4, 0.4,
            "📎 Code: github.com/zhaocorey/M-DaQ", font_size=18, color=C_DARK_BLUE)
add_textbox(slide, COL3_X + 0.2, y3 + 0.45, COL_W - 0.4, 0.4,
            "📄 Paper: arXiv:2509.15549", font_size=18, color=C_DARK_BLUE)
add_textbox(slide, COL3_X + 0.2, y3 + 0.8, COL_W - 0.4, 0.4,
            "📧 Contact: liuyilun3@huawei.com", font_size=18, color=C_DARK_BLUE)

# ── Footer ──
add_textbox(slide, 0, POSTER_H_INCHES - 0.6, POSTER_W_INCHES, 0.5,
            "SIGIR 2026 · Melbourne, VIC, Australia · July 20–24, 2026",
            font_size=16, color=C_GRAY, alignment=PP_ALIGN.CENTER)

# ── Save ──
output_path = os.path.join(OUTPUT_DIR, "M-DaQ_SIGIR2026_poster.pptx")
prs.save(output_path)
print(f"\n✅ Poster saved: {output_path}")
print(f"   Size: A0 Portrait ({POSTER_W_INCHES} × {POSTER_H_INCHES} inches)")
print(f"   Layout: 3 columns")
print(f"   Figures used: {sum(1 for f in [pipeline_fig, winrate_fig, sah_fig, corsica_fig] if f)}")
