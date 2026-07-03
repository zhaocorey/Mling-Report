#!/usr/bin/env python3
"""M-DaQ Poster v3 for SIGIR 2026.
A0 Portrait, 3-column. Layout fixes:
  - Zone-based vertical spacing → content fills full page height
  - Right column occlusion resolved (proper figure sizing)
  - Company logos (Huawei + USTC) + conference badge in title bar
  - Updated contact: DOI 10.1145/3805712.3809946, arXiv:2509.15549,
    zhaochunguang6@huawei.com, +86-15501192353
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Cm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ── Config ──
PAPER_PDF = "paper.pdf"
OUTPUT_DIR = "_output"
W = 33.1   # A0 width inches
H = 46.8   # A0 height inches

# Colors
C_DARK   = RGBColor(0x1B, 0x3A, 0x5C)
C_ACCENT = RGBColor(0xE8, 0x79, 0x2B)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_BLACK  = RGBColor(0x22, 0x22, 0x22)
C_GRAY   = RGBColor(0x66, 0x66, 0x66)
C_LGRAY  = RGBColor(0x99, 0x99, 0x99)
C_BG     = RGBColor(0xF0, 0xF4, 0xF8)
C_LBLUE  = RGBColor(0xCC, 0xDD, 0xEE)

LOGO_DIR = "logos"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ── Extract figures from PDF ──
import fitz
print("📄 Extracting figures...")
doc = fitz.open(PAPER_PDF)
figs = {}
for pi in range(len(doc)):
    page = doc[pi]
    for ii, info in enumerate(page.get_images(full=True)):
        base = doc.extract_image(info[0])
        if base and base["width"] > 200 and base["height"] > 150:
            key = f"fig_p{pi}_{ii}"
            path = os.path.join(OUTPUT_DIR, f"{key}.png")
            with open(path, "wb") as f:
                f.write(base["image"])
            figs[key] = {"path": path, "w": base["width"], "h": base["height"]}
            print(f"  ✅ {key}: {base['width']}×{base['height']}")
doc.close()

# ── Presentation ──
print("\n🎨 Generating poster...")
prs = Presentation()
prs.slide_width  = Inches(W)
prs.slide_height = Inches(H)
sl = prs.slides.add_slide(prs.slide_layouts[6])

# ── Helpers ──
def sc(run, rgb):
    """Set run color via XML."""
    hv = str(rgb)
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn('a:solidFill')):
        rPr.remove(old)
    sf = etree.SubElement(rPr, qn('a:solidFill'))
    sc = etree.SubElement(sf, qn('a:srgbClr'))
    sc.set('val', hv)

def tb(left, top, w, h, text, sz=24, bold=False, color=C_BLACK,
       align=PP_ALIGN.LEFT, font="Arial", fill=None, anchor=None):
    box = sl.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        tf.paragraphs[0].alignment = align
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(sz)
    run.font.bold = bold
    run.font.name = font
    sc(run, color)
    p.alignment = align
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    return box

def multi(left, top, w, h, paras, fill=None):
    """paras = [(text, size, bold, color, align, space_after), ...]"""
    box = sl.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (txt, sz, b, col, al, sp) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if txt:
            run = p.add_run()
            run.text = txt
            run.font.size = Pt(sz)
            run.font.bold = b
            run.font.name = "Arial"
            sc(run, col)
        p.alignment = al
        p.space_after = Pt(sp)
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    return box

def blt(left, top, w, h, items, sz=22, color=C_BLACK, sp=Pt(8)):
    box = sl.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"• {item}"
        run.font.size = Pt(sz)
        run.font.name = "Arial"
        sc(run, color)
        p.space_after = sp
    return box

def hdr(x, y, text, num=None):
    label = f"{num}  {text}" if num else text
    tb(x, y, CW, 0.55, label, sz=28, bold=True, color=C_DARK)
    line = sl.shapes.add_shape(1, Inches(x), Inches(y + 0.55),
                                Inches(CW), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = C_ACCENT
    line.line.fill.background()
    return y + 0.8

def img(path, left, top, width=None, height=None):
    if not os.path.exists(path):
        return 0
    kw = {"image_file": path, "left": Inches(left), "top": Inches(top)}
    if width:  kw["width"]  = Inches(width)
    if height: kw["height"] = Inches(height)
    sl.shapes.add_picture(**kw)
    if width and path in [v["path"] for v in figs.values()]:
        for v in figs.values():
            if v["path"] == path:
                return min(v["h"] * width / v["w"], 99)
    return 0

# ── Layout: margins, columns ──
MARGIN  = 0.8
GAP     = 0.5
CW      = (W - 2 * MARGIN - 2 * GAP) / 3   # col width ≈ 10.17"
C1      = MARGIN
C2      = MARGIN + CW + GAP
C3      = MARGIN + 2 * (CW + GAP)

# ═══════════════════════════════════════
# TITLE BAR  (0 → 3.8")
# ═══════════════════════════════════════
TH = 3.8
bg = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(W), Inches(TH))
bg.fill.solid()
bg.fill.fore_color.rgb = C_DARK
bg.line.fill.background()

# Title
tb(0.8, 0.25, W - 1.6, 1.5,
   "M-DaQ: Retrieving Samples with Multilingual Diversity\nand Quality for Instruction Fine-Tuning Datasets",
   sz=44, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

# Authors
tb(0.8, 1.75, W - 1.6, 0.5,
   "Chunguang Zhao¹ · Yilun Liu¹ · Pufan Zeng² · Yuanchang Luo¹ · Shimin Tao¹ · Minggui He¹ · Weibin Meng¹ · Song Xu² · Chen Liu¹ · Hongxia Ma¹ · Li Zhang¹ · Boxing Chen¹ · Daimeng Wei¹",
   sz=18, color=C_WHITE, align=PP_ALIGN.CENTER)

# Affiliations
tb(0.8, 2.2, W - 1.6, 0.4,
   "¹ Huawei Technologies Ltd.    ² University of Science and Technology of China",
   sz=17, color=RGBColor(0xBB, 0xCC, 0xDD), align=PP_ALIGN.CENTER)

# Logos row
ly = 2.75
lh = 0.75
# Huawei logo (left)
hlogo = os.path.join(LOGO_DIR, "huawei_logo.png")
if os.path.exists(hlogo):
    img(hlogo, 1.5, ly, height=lh)
# USTC logo (right)
ulogo = os.path.join(LOGO_DIR, "ustc_logo.png")
if os.path.exists(ulogo):
    img(ulogo, W - 5.0, ly, height=lh)
# SIGIR badge (center)
tb(W/2 - 4, ly + 0.1, 8, 0.6,
   "SIGIR 2026  ·  Melbourne, Australia  ·  July 20–24",
   sz=22, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)

YS = TH + 0.5        # content start Y = 4.3
FY = H - 0.7         # footer Y = 46.1
AVAIL = FY - YS      # available height = 41.8

# ═══════════════════════════════════════
# ZONE-BASED LAYOUT
# Each column gets the full available height.
# We pre-calculate zones and space content within them.
# ═══════════════════════════════════════

# ── COLUMN 1: Motivation (zone 1) + Method + Figure (zone 2) ──
# Zone split: 35% motivation, 65% method+figure
Z1_H = AVAIL * 0.35   # ~14.6"
Z2_H = AVAIL * 0.65   # ~27.2"

# Zone 1: Motivation
y1 = YS
y1 = hdr(C1, y1, "MOTIVATION", "1")

blt(C1, y1, CW, 6.0, [
    "Multilingual IFT data is scarce, skewed toward English, and lacks systematic curation",
    "Llama-3 IFT dataset: only 3.01% multilingual samples",
    "No language-agnostic quality scoring for multilingual data",
    "Data diversity studied only in English settings",
    "Superficial Alignment Hypothesis (SAH) unverified in multilingual contexts",
], sz=22, sp=Pt(16))
y1 += 7.5

# Challenges sub-section
tb(C1, y1, CW, 0.4, "Three Challenges Addressed:", sz=22, bold=True, color=C_DARK)
y1 += 0.6
for cid, ch, sol in [
    ("C1", "No extensible quality scoring for multilingual IFT", "→ QSM with triplet loss"),
    ("C2", "Diversity selection limited to English", "→ DAS inspired by MMR"),
    ("C3", "SAH unverified in multilingual settings", "→ 1K→52K study, 8 languages"),
]:
    tb(C1, y1, 0.6, 0.4, cid, sz=20, bold=True, color=C_ACCENT)
    multi(C1 + 0.6, y1, CW - 0.6, 0.8, [
        (ch, 19, False, C_BLACK, PP_ALIGN.LEFT, 2),
        (sol, 18, True, C_ACCENT, PP_ALIGN.LEFT, 4),
    ])
    y1 += 1.1

# Pad remaining zone 1 space
y1 = YS + Z1_H + 0.3

# Zone 2: Method + Pipeline Figure
y1 = hdr(C1, y1, "METHOD: M-DaQ FRAMEWORK", "2")

tb(C1, y1, CW, 0.4, "Stage 1: Quality Scoring Model (QSM)",
   sz=22, bold=True, color=C_ACCENT)
y1 += 0.5
blt(C1, y1, CW, 3.0, [
    "Fine-tuned on ~2.3K expert-revised samples × 18 languages (from MIDB)",
    "Triplet loss: instruction → positive (expert-revised) vs negative (original/MT)",
    "Language-agnostic quality signal from embedding space",
], sz=20, sp=Pt(12))
y1 += 3.5

tb(C1, y1, CW, 0.4, "Stage 2: Diversity-Aware Selection (DAS)",
   sz=22, bold=True, color=C_ACCENT)
y1 += 0.5
blt(C1, y1, CW, 4.5, [
    "MMR-inspired: Quality ≈ Relevance, Diversity ≈ Novelty",
    "Two-stage pipeline reduces O(n²) → O(n log n)",
    "  Stage A: Select top-n by QSM score (quality subset)",
    "  Stage B: Greedily augment from uncovered clusters (diversity subset)",
    "Ratio n_quality : n_diversity = 6:1 (empirically tuned)",
], sz=20, sp=Pt(12))
y1 += 5.0

# Pipeline figure — fill remaining zone 2 space
pipe = figs.get("fig_p1_1")
if pipe:
    remaining = (YS + Z1_H + Z2_H) - y1 - 0.6  # space left
    fig_h = min(remaining, pipe["h"] * CW / pipe["w"])
    fig_h = max(fig_h, 3.0)  # minimum 3"
    img(pipe["path"], C1, y1, width=CW)
    tb(C1, y1 + fig_h + 0.05, CW, 0.4,
       "Figure: M-DaQ two-stage pipeline (QSM + DAS)",
       sz=16, color=C_GRAY, align=PP_ALIGN.CENTER)

# ── COLUMN 2: Results (zone 1) + Human Eval (zone 2) + Cultural (zone 3) ──
# Zone split: 40% results, 30% human eval, 30% cultural
Z2A = AVAIL * 0.40   # ~16.7"
Z2B = AVAIL * 0.30   # ~12.5"
Z2C = AVAIL * 0.30   # ~12.5"

y2 = YS
y2 = hdr(C2, y2, "KEY RESULTS", "3")

# Stats highlight card
card = sl.shapes.add_shape(1, Inches(C2), Inches(y2), Inches(CW), Inches(1.4))
card.fill.solid()
card.fill.fore_color.rgb = C_BG
card.line.fill.background()

multi(C2 + 0.2, y2 + 0.1, CW - 0.4, 1.2, [
    ("LLM-as-Judge: Avg Win Rate (M-DaQ vs Vanilla)", 20, True, C_DARK, PP_ALIGN.CENTER, 10),
    ("Alpaca-Eval: 60.2%  |  MT-Bench R1: 62.6%  |  MT-Bench R2: 62.9%",
     20, True, C_ACCENT, PP_ALIGN.CENTER, 4),
])
y2 += 1.7

# Win rate figure (hero)
wrf = figs.get("fig_p1_0")  # 2491×1266
if wrf:
    zone_end = YS + Z2A - 1.0  # leave 1" for caption
    max_fig_h = zone_end - y2
    fig_h = min(max_fig_h, wrf["h"] * CW / wrf["w"])
    fig_h = max(fig_h, 4.0)
    img(wrf["path"], C2, y2, width=CW)
    tb(C2, y2 + fig_h + 0.05, CW, 0.5,
       "Cross-lingual win rates across 18 languages (Alpaca-Eval + MT-Bench)",
       sz=16, color=C_GRAY, align=PP_ALIGN.CENTER)
    y2 = YS + Z2A

# Zone 2B: Key observations + Human eval table
y2 += 0.3
tb(C2, y2, CW, 0.4, "Key Observations:", sz=22, bold=True, color=C_DARK)
y2 += 0.5
blt(C2, y2, CW, 3.0, [
    "Consistent improvement across all 18 languages",
    "Gains MORE pronounced in low-resource languages (Tagalog, Malay)",
    "M-DaQ effectively mitigates linguistic imbalance and data scarcity",
], sz=20, sp=Pt(12))
y2 += 3.3

# Human eval table
tb(C2, y2, CW, 0.4,
   "Human Evaluation (6 Langs, 900 Samples, 58 Person-Hours):",
   sz=19, bold=True, color=C_DARK)
y2 += 0.5

rows = [
    ("Language", "Alpaca-Eval", "MT-Bench R1", "MT-Bench R2"),
    ("Japanese",   "86%", "84%", "80%"),
    ("Korean",     "94%", "94%", "94%"),
    ("Russian",    "70%", "86%", "84%"),
    ("Portuguese", "80%", "78%", "84%"),
    ("Greek",      "58%", "76%", "82%"),
    ("French",     "80%", "92%", "88%"),
    ("Average",    "78.0%", "85.0%", "85.3%"),
]
nrows = len(rows)
row_h = 0.40
tbl_shape = sl.shapes.add_table(nrows, 4,
    Inches(C2), Inches(y2), Inches(CW), Inches(row_h * nrows))
tbl = tbl_shape.table

# Set column widths
col_widths = [CW * 0.30, CW * 0.24, CW * 0.23, CW * 0.23]
for ci, cw in enumerate(col_widths):
    tbl.columns[ci].width = Inches(cw)

for ci, h in enumerate(rows[0]):
    cell = tbl.cell(0, ci)
    cell.text = h
    for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.name = "Arial"
            sc(r, C_WHITE)
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_DARK

for ri, row in enumerate(rows[1:], 1):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.text = val
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(15)
                r.font.name = "Arial"
                is_avg = (ri == nrows - 1)
                r.font.bold = is_avg
                sc(r, C_DARK if (ci == 0 or is_avg) else C_BLACK)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_BG if ri % 2 == 0 else C_WHITE

y2 = YS + Z2A + Z2B + 0.5

# Zone 2C: Cultural Localization
tb(C2, y2, CW, 0.4, "Cultural Localization Example:",
   sz=22, bold=True, color=C_DARK)
y2 += 0.5

cor = figs.get("fig_p4_0")  # 828×414
if cor:
    zone_end = FY - 0.5
    max_fig_h = zone_end - y2 - 0.6
    fig_h = min(max_fig_h, cor["h"] * CW / cor["w"])
    fig_h = max(fig_h, 3.0)
    img(cor["path"], C2, y2, width=CW)
    tb(C2, y2 + fig_h + 0.1, CW, 0.8,
       "French travel query about Corsica → M-DaQ references Piana, figatelli, local cultural conventions",
       sz=16, color=C_GRAY, align=PP_ALIGN.CENTER)

# ── COLUMN 3: SAH (zone 1) + Conclusion (zone 2) + Contact (zone 3) ──
# Zone split: 48% SAH, 30% Conclusion, 22% Contact
Z3A = AVAIL * 0.48   # ~20.1"
Z3B = AVAIL * 0.30   # ~12.5"
Z3C = AVAIL * 0.22   # ~9.2"

y3 = YS
y3 = hdr(C3, y3, "SAH IN MULTILINGUAL SETTINGS", "4")

tb(C3, y3, CW, 0.7,
   "First systematic investigation of the Superficial Alignment Hypothesis across languages.",
   sz=20)
y3 += 0.8

# SAH figure
sah = figs.get("fig_p4_1")  # 1200×840
if sah:
    max_fig_h = Z3A * 0.45
    fig_h = min(max_fig_h, sah["h"] * CW / sah["w"])
    fig_h = max(fig_h, 4.0)
    img(sah["path"], C3, y3, width=CW)
    tb(C3, y3 + fig_h + 0.05, CW, 0.4,
       "Win rate vs IFT dataset scale (1K–52K), 8 languages",
       sz=16, color=C_GRAY, align=PP_ALIGN.CENTER)
    y3 += fig_h + 0.8

# SAH findings
tb(C3, y3, CW, 0.4, "Key Findings:", sz=22, bold=True, color=C_DARK)
y3 += 0.5
blt(C3, y3, CW, 6.0, [
    "✅  Diminishing returns: 1K curated > 52K unfiltered",
    "     — SAH holds multilingually",
    "     10K → −10.1% avg win rate vs 1K baseline",
    "     52K → −6.2% avg win rate vs 1K baseline",
    "⚠️  Language-dependent sensitivity:",
    "     Arabic shows flatter curve than French",
    "     (lower pretraining readiness)",
    "💡  A few thousand high-quality samples suffice",
    "     for effective multilingual alignment",
], sz=20, sp=Pt(10))

# Zone 2: Conclusion
y3 = YS + Z3A + 0.5
y3 = hdr(C3, y3, "CONCLUSION", "5")
blt(C3, y3, CW, 7.0, [
    "M-DaQ = QSM (quality) + DAS (diversity):",
    "  compact, high-fidelity multilingual IFT subsets",
    "60%+ avg win rate across 18 languages",
    "  on Alpaca-Eval & MT-Bench",
    "Human eval confirms gains in cultural relevance,",
    "  contextual appropriateness, instruction-following",
    "First empirical validation of SAH",
    "  in multilingual settings",
    "Code released: github.com/zhaocorey/M-DaQ",
], sz=20, sp=Pt(10))

# Zone 3: Contact card
y3 = YS + Z3A + Z3B + 0.8
card_h = min(Z3C - 0.5, 4.5)
card = sl.shapes.add_shape(1, Inches(C3), Inches(y3),
                            Inches(CW), Inches(card_h))
card.fill.solid()
card.fill.fore_color.rgb = C_BG
card.line.fill.background()

multi(C3 + 0.3, y3 + 0.2, CW - 0.6, card_h - 0.3, [
    ("📎 Code & Resources", 22, True, C_DARK, PP_ALIGN.LEFT, 10),
    ("GitHub: github.com/zhaocorey/M-DaQ", 18, False, C_BLACK, PP_ALIGN.LEFT, 12),
    ("📄 Paper", 22, True, C_DARK, PP_ALIGN.LEFT, 10),
    ("DOI: 10.1145/3805712.3809946", 18, False, C_BLACK, PP_ALIGN.LEFT, 4),
    ("arXiv: 2509.15549", 18, False, C_BLACK, PP_ALIGN.LEFT, 12),
    ("📧 Contact", 22, True, C_DARK, PP_ALIGN.LEFT, 10),
    ("zhaochunguang6@huawei.com", 18, False, C_BLACK, PP_ALIGN.LEFT, 4),
    ("+86-15501192353", 18, False, C_BLACK, PP_ALIGN.LEFT, 4),
])

# ── Footer ──
tb(0, FY, W, 0.5,
   "SIGIR 2026  ·  Melbourne, VIC, Australia  ·  July 20–24, 2026",
   sz=16, color=C_GRAY, align=PP_ALIGN.CENTER)

# ── Save ──
out_pptx = os.path.join(OUTPUT_DIR, "M-DaQ_poster_v3.pptx")
prs.save(out_pptx)
print(f"\n✅ PPTX: {out_pptx}")

# PDF + PNG via LibreOffice
import subprocess
try:
    subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf",
                     "--outdir", OUTPUT_DIR, out_pptx],
                    timeout=60, capture_output=True)
    out_pdf = out_pptx.replace(".pptx", ".pdf")
    if os.path.exists(out_pdf):
        print(f"✅ PDF:  {out_pdf}")
        subprocess.run(["libreoffice", "--headless", "--convert-to", "png",
                         "--outdir", OUTPUT_DIR, out_pdf],
                        timeout=60, capture_output=True)
        out_png = out_pptx.replace(".pptx", "_preview.png")
        raw_png = out_pptx.replace(".pptx", ".png")
        if os.path.exists(raw_png):
            os.rename(raw_png, out_png)
            print(f"✅ PNG:  {out_png}")
except Exception as e:
    print(f"⚠️  LibreOffice: {e}")

print(f"\n   A0 Portrait, 3 columns, zone-based layout")
print(f"   Figures: {len(figs)}")
